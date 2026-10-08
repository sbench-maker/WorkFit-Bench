from __future__ import annotations

import csv
import json
import os
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader


RESULTS = Path(os.environ.get("RESULTS_DIR", "/root/results"))
DATA = Path(os.environ.get("DATA_DIR", "/root/data"))
XLSX = RESULTS / "close-packet-2026-04.xlsx"
PDF = RESULTS / "close-packet-2026-04.pdf"


def norm(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().replace("&", " and ").replace("quick books", "quickbooks")
    text = text.replace("—", "-").replace("–", "-").replace("_", " ")
    return re.sub(r"[^a-z0-9%]+", " ", text).strip()


def identifier(value: object) -> str:
    raw = str(value or "").strip()
    if not raw or re.fullmatch(r"[-—–]+", raw) or norm(raw) in {"na", "n a", "none", "not applicable"}:
        return ""
    return raw.upper()


def parse_money(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return int(round(float(value) * 100))
    text = str(value).strip()
    if not text or norm(text) in {"na", "n a", "none", "not applicable"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if not cleaned or cleaned in {"-", "."}:
        return None
    number = Decimal(cleaned)
    if negative:
        number = -abs(number)
    return int(round(number * 100))


def parse_ratio(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = float(value)
    else:
        match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
        if not match:
            return None
        number = float(match.group())
        if "%" in str(value):
            number /= 100
    return number / 100 if abs(number) > 1.5 else number


def read_csv(name: str) -> list[dict[str, str]]:
    with (DATA / name).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


@lru_cache(maxsize=1)
def packet_state() -> dict:
    state: dict = {"errors": []}
    if not XLSX.is_file():
        state["errors"].append("close-packet-2026-04.xlsx is missing")
    else:
        try:
            state["workbook"] = load_workbook(XLSX, data_only=True, read_only=False)
        except Exception as exc:
            state["errors"].append(f"the workbook cannot be opened: {exc}")
    if not PDF.is_file():
        state["errors"].append("close-packet-2026-04.pdf is missing")
    else:
        try:
            reader = PdfReader(str(PDF))
            state["pdf_reader"] = reader
            state["pdf_text"] = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            state["errors"].append(f"the PDF cannot be opened: {exc}")
    return state


ALIASES = {
    "processor": ["processor", "payment processor"],
    "settlement_id": ["settlement id", "payout id", "processor settlement id"],
    "settlement_date": ["settlement date", "payout date", "processor date"],
    "processor_amount": ["processor net", "net amount", "settlement amount", "payout amount"],
    "qb_id": ["quickbooks transaction id", "qb transaction id", "quickbooks id", "qb id"],
    "qb_date": ["quickbooks date", "qb date", "deposit date"],
    "qb_amount": ["quickbooks amount", "qb amount", "deposit amount"],
    "variance": ["variance", "difference"],
    "status": ["status", "match status", "reconciliation status"],
    "flag_id": ["flag id", "issue id", "exception id"],
    "issue_type": ["issue type", "flag type", "exception type"],
    "paired_id": ["paired transaction id", "paired id", "possible duplicate id", "duplicate transaction id"],
    "owner_decision": ["owner decision", "triage decision", "decision", "disposition"],
    "approval_status": ["approval status", "triage status", "review status"],
    "final_category": ["final category", "approved category", "reporting category"],
    "excluded_id": ["excluded transaction id", "excluded id", "duplicate excluded"],
    "action": ["recommended action", "next action", "action"],
    "line_item": ["line item", "account category", "p and l line", "category", "line"],
    "march": ["march 2026", "mar 2026", "2026 03", "prior month", "march", "mar"],
    "april": ["april 2026", "apr 2026", "2026 04", "current month", "april", "apr"],
    "account_number": ["account number", "account no", "account code", "number"],
    "account_name": ["account name", "account"],
    "ending_balance": ["ending balance", "balance usd", "closing balance", "balance"],
}


def header_field(value: object, candidates: set[str]) -> str | None:
    value_norm = norm(value)
    if not value_norm:
        return None
    matches: list[tuple[int, int, str]] = []
    for field in candidates:
        aliases = [norm(alias) for alias in ALIASES[field]]
        for alias in aliases:
            if value_norm == alias:
                matches.append((3, len(alias), field))
            elif value_norm.startswith(alias + " "):
                matches.append((2, len(alias), field))
            elif alias in value_norm:
                matches.append((1, len(alias), field))
    return max(matches)[2] if matches else None


def find_table(
    required: set[str], optional: set[str] | None = None, *, skip_missing: bool = True
) -> tuple[object, dict[str, int], int]:
    state = packet_state()
    if state["errors"] or "workbook" not in state:
        pytest.skip("packet is unreadable; root artifact failure is reported separately")
    candidates = required | (optional or set())
    best = None
    for ws in state["workbook"].worksheets:
        for row_idx in range(1, min(ws.max_row, 40) + 1):
            mapping: dict[str, int] = {}
            for col_idx in range(1, min(ws.max_column, 30) + 1):
                field = header_field(ws.cell(row_idx, col_idx).value, candidates)
                if field and field not in mapping:
                    mapping[field] = col_idx
            if required.issubset(mapping):
                score = len(mapping)
                if best is None or score > best[0]:
                    best = (score, ws, mapping, row_idx)
    if best is None:
        message = f"no usable workbook table exposes required semantic fields {sorted(required)}"
        if skip_missing:
            pytest.skip(message)
        raise AssertionError(message)
    return best[1], best[2], best[3]


def table_rows(ws, mapping: dict[str, int], header_row: int) -> list[dict[str, object]]:
    output = []
    blank_run = 0
    for row_idx in range(header_row + 1, ws.max_row + 1):
        row = {field: ws.cell(row_idx, col).value for field, col in mapping.items()}
        if all(value is None or str(value).strip() == "" for value in row.values()):
            blank_run += 1
            if blank_run >= 3:
                break
            continue
        blank_run = 0
        output.append(row)
    return output


def expected_reconciliation() -> list[dict]:
    qb_rows = read_csv("qb_transactions.csv")
    settlements = read_csv("processor_settlements.csv")
    deposits = [row for row in qb_rows if row["date"].startswith("2026-04") and row["transaction_type"] == "Deposit"]
    used: set[str] = set()
    output = []
    for settlement in settlements:
        sdate = date.fromisoformat(settlement["settlement_date"])
        amount = parse_money(settlement["net_amount_usd"])
        candidates = [
            row
            for row in deposits
            if row["transaction_id"] not in used and abs((date.fromisoformat(row["date"]) - sdate).days) <= 2
        ]
        exact = [row for row in candidates if parse_money(row["amount_usd"]) == amount]
        referenced = [row for row in candidates if row["source_reference"] == settlement["settlement_id"]]
        match = exact[0] if exact else (referenced[0] if referenced else None)
        if match:
            used.add(match["transaction_id"])
            qb_amount = parse_money(match["amount_usd"])
            variance = qb_amount - amount
            status = "matched" if variance == 0 else "variance"
            qb_id = match["transaction_id"]
        else:
            qb_amount, variance, qb_id, status = None, None, "", "unmatched_processor"
        output.append(
            {
                "settlement_id": settlement["settlement_id"],
                "qb_id": qb_id,
                "processor_amount": amount,
                "qb_amount": qb_amount,
                "variance": variance,
                "status": status,
            }
        )
    for row in deposits:
        if row["transaction_id"] not in used:
            output.append(
                {
                    "settlement_id": "",
                    "qb_id": row["transaction_id"],
                    "processor_amount": None,
                    "qb_amount": parse_money(row["amount_usd"]),
                    "variance": None,
                    "status": "unmatched_qb",
                }
            )
    return output


def status_key(value: object) -> str:
    text = norm(value)
    if "variance" in text or "amount difference" in text or "mismatch" in text:
        return "variance"
    if "unmatched" in text or "missing" in text or "no match" in text:
        if "processor" in text or "settlement" in text or "payout" in text:
            return "unmatched_processor"
        if "quickbooks" in text or re.search(r"\bqb\b", text) or "deposit" in text:
            return "unmatched_qb"
    if text in {"matched", "match", "reconciled", "exact match", "cleared"} or "exact match" in text:
        return "matched"
    return text


def issue_key(value: object) -> str:
    text = norm(value)
    if "uncategor" in text or "no category" in text:
        return "uncategorized"
    if "duplicate" in text:
        return "possible_duplicate"
    if "receipt" in text or "attachment" in text:
        return "missing_receipt"
    return text.replace(" ", "_")


def test_packet_artifacts_are_usable() -> None:
    state = packet_state()
    assert not state["errors"], "; ".join(state["errors"])
    workbook = state["workbook"]
    assert len(workbook.worksheets) >= 4, "the workbook does not expose the four close schedules described in the close notes"
    assert len(state["pdf_reader"].pages) == 1, "the requested close summary must remain one page"
    identity = norm(" ".join(str(cell.value or "") for ws in workbook.worksheets for row in ws.iter_rows() for cell in row))
    pdf_text = norm(state["pdf_text"])
    for text, artifact in ((identity, "workbook"), (pdf_text, "PDF")):
        assert "harbor and pine studio" in text, f"the {artifact} does not identify the company"
        assert "april 2026" in text, f"the {artifact} does not identify the close period"
    # Logical sections are located by headers, not exact sheet names.
    find_table(
        {"settlement_id", "qb_id", "status"},
        {"processor_amount", "qb_amount", "variance"},
        skip_missing=False,
    )
    find_table(
        {"issue_type", "qb_id", "owner_decision", "approval_status"},
        {"paired_id", "final_category", "action"},
        skip_missing=False,
    )
    find_table({"line_item", "march", "april"}, skip_missing=False)
    find_table({"account_name", "ending_balance"}, {"account_number"}, skip_missing=False)


def test_reconciliation_scope_and_classification() -> None:
    ws, mapping, header = find_table(
        {"settlement_id", "qb_id", "status"},
        {"processor_amount", "qb_amount", "variance", "processor", "settlement_date", "qb_date"},
    )
    actual_rows = table_rows(ws, mapping, header)
    actual: dict[tuple[str, str], dict] = {}
    duplicates = []
    for row in actual_rows:
        settlement_id = identifier(row.get("settlement_id"))
        qb_id = identifier(row.get("qb_id"))
        if not settlement_id and not qb_id:
            continue
        key = (settlement_id, qb_id)
        if key in actual:
            duplicates.append(key)
        actual[key] = row
    expected = {(row["settlement_id"].upper(), row["qb_id"].upper()): row for row in expected_reconciliation()}
    assert not duplicates, f"reconciliation rows are duplicated for {duplicates[:5]}"
    missing = sorted(set(expected) - set(actual))
    extra = sorted(set(actual) - set(expected))
    assert not missing and not extra, f"reconciliation scope differs from the frozen exports; missing={missing[:5]}, extra={extra[:5]}"
    errors = []
    for key, expected_row in expected.items():
        row = actual[key]
        if status_key(row.get("status")) != expected_row["status"]:
            errors.append(f"{key}: status {row.get('status')!r}")
            continue
        for field in ("processor_amount", "qb_amount"):
            if field in mapping and parse_money(row.get(field)) != expected_row[field]:
                errors.append(f"{key}: {field} is {row.get(field)!r}, expected cents {expected_row[field]}")
        if "variance" in mapping and expected_row["variance"] is not None:
            observed = parse_money(row.get("variance"))
            if observed is None or abs(observed) != abs(expected_row["variance"]):
                errors.append(f"{key}: variance is {row.get('variance')!r}, expected absolute cents {abs(expected_row['variance'])}")
    assert not errors, "material reconciliation errors: " + "; ".join(errors[:8])


def test_unavailable_processor_disclosure() -> None:
    state = packet_state()
    if state["errors"]:
        pytest.skip("packet is unreadable; root artifact failure is reported separately")
    workbook_text = " ".join(str(cell.value or "") for ws in state["workbook"].worksheets for row in ws.iter_rows() for cell in row)
    combined = norm(workbook_text + " " + state["pdf_text"])
    assert "paypal" in combined and "unavailable" in combined, "the packet does not disclose the unavailable PayPal feed"
    assert any(token in combined for token in ("skipped", "not connected", "not included", "excluded from reconciliation")), (
        "the PayPal disclosure does not explain that its settlements were not reconciled"
    )


def test_flagged_schedule_and_owner_triage() -> None:
    ws, mapping, header = find_table(
        {"issue_type", "qb_id", "owner_decision", "approval_status"},
        {"flag_id", "paired_id", "final_category", "excluded_id", "action"},
    )
    actual_rows = table_rows(ws, mapping, header)
    actual: dict[tuple[str, str, str], dict] = {}
    for row in actual_rows:
        issue = issue_key(row.get("issue_type"))
        qb_id = identifier(row.get("qb_id"))
        pair = identifier(row.get("paired_id"))
        if not issue or not qb_id:
            continue
        if issue == "possible_duplicate" and pair:
            qb_id, pair = sorted((qb_id, pair))
        actual[(issue, qb_id, pair)] = row

    expected_rows = read_csv("owner_triage.csv")
    expected = {}
    for item in expected_rows:
        qb_id = item["qb_transaction_id"].upper()
        pair = item["paired_transaction_id"].upper()
        if item["issue_type"] == "possible_duplicate":
            qb_id, pair = sorted((qb_id, pair))
        expected[(item["issue_type"], qb_id, pair)] = item
    assert set(actual) == set(expected), (
        f"flagged scope does not match the review rules and owner triage; missing={sorted(set(expected)-set(actual))[:5]}, "
        f"extra={sorted(set(actual)-set(expected))[:5]}"
    )

    errors = []
    for key, item in expected.items():
        row = actual[key]
        row_text = norm(" ".join(str(value or "") for value in row.values()))
        decision = item["owner_decision"]
        if decision == "categorize":
            ok = "categor" in row_text and norm(item["final_category"]) in row_text
        elif decision == "exclude_confirmed":
            ok = "exclud" in row_text and norm(item["excluded_transaction_id"]) in row_text
        elif decision == "keep_both":
            ok = "keep both" in row_text or "separate" in row_text
        elif decision == "needs_follow_up":
            ok = ("follow" in row_text or "obtain" in row_text or "open" in row_text) and "open" in row_text
        elif decision == "receipt_attached_after_review":
            ok = "receipt" in row_text and ("attach" in row_text or "located" in row_text) and "approved" in row_text
        else:
            ok = "receipt" in row_text and ("request" in row_text or "obtain" in row_text) and "open" in row_text
        if not ok:
            errors.append(f"{key}: owner disposition is not faithfully carried forward")
    assert not errors, "; ".join(errors[:8])


def expected_pnl() -> dict[str, dict[str, int]]:
    triage = read_csv("owner_triage.csv")
    excluded = {
        row["excluded_transaction_id"]
        for row in triage
        if row["approval_status"] == "approved" and row["owner_decision"] == "exclude_confirmed"
    }
    overrides = {
        row["qb_transaction_id"]: row["final_category"]
        for row in triage
        if row["approval_status"] == "approved" and row["owner_decision"] == "categorize"
    }
    category: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    sections: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for row in read_csv("qb_transactions.csv"):
        month = row["date"][:7]
        if month not in {"2026-03", "2026-04"} or not row["pl_section"] or row["transaction_id"] in excluded:
            continue
        name = overrides.get(row["transaction_id"], row["category"] or "Uncategorized")
        value = parse_money(row["amount_usd"])
        category[month][name] += value
        sections[month][row["pl_section"]] += value
    output = {month: dict(values) for month, values in category.items()}
    for month in ("2026-03", "2026-04"):
        revenue = sections[month]["Revenue"]
        cogs = sections[month]["Cost of Goods Sold"]
        opex = sections[month]["Operating Expense"]
        output[month].update(
            {
                "Total Revenue": revenue,
                "Total Cost of Goods Sold": cogs,
                "Gross Profit": revenue + cogs,
                "Total Operating Expenses": opex,
                "Net Income": revenue + cogs + opex,
                "Gross Margin": round((revenue + cogs) / revenue * 100_000_000),
            }
        )
    return output


P_AND_L_ALIASES = {
    "total revenue": "Total Revenue",
    "revenue total": "Total Revenue",
    "total cost of goods sold": "Total Cost of Goods Sold",
    "total cogs": "Total Cost of Goods Sold",
    "cost of goods sold": "Total Cost of Goods Sold",
    "gross profit": "Gross Profit",
    "gross margin": "Gross Margin",
    "total operating expenses": "Total Operating Expenses",
    "total operating expense": "Total Operating Expenses",
    "total opex": "Total Operating Expenses",
    "net income": "Net Income",
    "net profit": "Net Income",
}


def canonical_pnl_label(value: object, expected_labels: set[str]) -> str | None:
    text = norm(value)
    if text in P_AND_L_ALIASES:
        return P_AND_L_ALIASES[text]
    for label in expected_labels:
        if text == norm(label):
            return label
    return None


def test_pnl_accuracy_and_adjustments() -> None:
    ws, mapping, header = find_table({"line_item", "march", "april"})
    rows = table_rows(ws, mapping, header)
    truth = expected_pnl()
    expected_labels = set(truth["2026-03"]) | set(truth["2026-04"])
    actual = {}
    for row in rows:
        label = canonical_pnl_label(row.get("line_item"), expected_labels)
        if label:
            actual[label] = row
    missing = sorted(expected_labels - set(actual))
    assert not missing, f"the March-versus-April P&L omits source-derived lines: {missing}"
    errors = []
    for label in sorted(expected_labels):
        if label == "Gross Margin":
            for field, month in (("march", "2026-03"), ("april", "2026-04")):
                observed = parse_ratio(actual[label].get(field))
                expected = truth[month][label] / 100_000_000
                if observed is None or abs(observed - expected) > 0.0006:
                    errors.append(f"{label} {month}: {actual[label].get(field)!r} vs {expected:.6f}")
        else:
            for field, month in (("march", "2026-03"), ("april", "2026-04")):
                observed = parse_money(actual[label].get(field))
                if observed != truth[month].get(label, 0):
                    errors.append(f"{label} {month}: {observed} cents vs {truth[month].get(label, 0)}")
    assert not errors, "P&L values do not reflect the frozen books and approved triage: " + "; ".join(errors[:8])


def test_trial_balance_accuracy() -> None:
    ws, mapping, header = find_table({"account_name", "ending_balance"}, {"account_number"})
    actual_rows = table_rows(ws, mapping, header)
    actual_by_number = {
        norm(row.get("account_number")): row for row in actual_rows if "account_number" in mapping and norm(row.get("account_number"))
    }
    actual_by_name = {norm(row.get("account_name")): row for row in actual_rows if norm(row.get("account_name"))}
    errors = []
    observed_total = 0
    for item in read_csv("trial_balance.csv"):
        row = actual_by_number.get(norm(item["account_number"])) or actual_by_name.get(norm(item["account_name"]))
        if row is None:
            errors.append(f"missing account {item['account_number']} {item['account_name']}")
            continue
        observed = parse_money(row.get("ending_balance"))
        expected = parse_money(item["ending_balance_usd"])
        if observed != expected:
            errors.append(f"{item['account_name']}: {observed} cents vs {expected}")
        observed_total += observed or 0
    assert not errors, "trial balance does not match the supplied account snapshot: " + "; ".join(errors[:8])
    assert observed_total == 0, f"the submitted trial-balance account rows do not balance; net cents={observed_total}"


def test_pdf_summary_consistency() -> None:
    state = packet_state()
    if state["errors"]:
        pytest.skip("packet is unreadable; root artifact failure is reported separately")
    text = state["pdf_text"].replace("−", "-").replace("–", "-").replace("—", "-")
    compact = re.sub(r"\s+", " ", text)
    assert re.search(r"\$\s*118[, ]?420(?:\.0+)?", compact), "the PDF does not show April revenue of $118,420"
    assert re.search(r"59\.0\s*%", compact), "the PDF does not show April gross margin rounded to 59.0%"
    assert re.search(r"\$\s*12[, ]?337(?:\.0+)?", compact), "the PDF does not show April net income rounded to $12,337"
    assert re.search(r"(?:open\s+recon(?:ciliation)?\s+gaps?|recon(?:ciliation)?\s+gaps?).{0,20}\b15\b", norm(compact)), (
        "the PDF does not report the 15 remaining reconciliation gaps"
    )
