from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook


OUTPUT = Path(os.environ.get("SUBMISSION_PATH", "/root/results/credit_cleaning_report.xlsx"))
DATA = Path(os.environ.get("TASK_DATA_DIR", "/root/data"))
FEATURE_RE = re.compile(r"\bi_[a-z0-9_]+\b", re.I)
DROP_WORDS = ("drop", "exclude", "remove", "reject", "fail", "剔除", "去除", "删除", "不保留")
KEEP_WORDS = ("keep", "retain", "include", "selected", "pass", "保留", "入模", "通过")
GATE_WORDS = {
    "zero_variance": ("zero variance", "zero_variance", "constant", "one distinct", "single unique", "零方差", "常量"),
    "missingness": ("missing", "缺失"),
    "iv": (" iv", "iv ", "iv_", "information value", "信息值"),
    "psi": ("psi", "stability", "unstable", "稳定"),
    "null_importance": ("null importance", "null_importance", "permut", "noise", "置换", "噪声"),
    "correlation": ("correlation", "correlated", "corr", "相关"),
}


def norm(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def numeric_values(row) -> list[float]:
    values = []
    for cell in row:
        if isinstance(cell, (int, float)) and not isinstance(cell, bool):
            values.append(float(cell))
        elif isinstance(cell, str):
            for match in re.findall(r"(?<![a-z_])[-+]?\d+(?:\.\d+)?%?", cell.lower()):
                try:
                    values.append(float(match.rstrip("%")) / (100 if match.endswith("%") else 1))
                except ValueError:
                    pass
    return values


@dataclass
class WorkbookView:
    error: str | None
    sheets: list[tuple[str, list[tuple]]]

    @property
    def rows(self):
        for title, rows in self.sheets:
            for row in rows:
                yield title, row

    def matching_rows(self, *needles: str, sheet_terms: tuple[str, ...] = ()):
        for title, row in self.rows:
            title_text = norm(title)
            row_text = " | ".join(norm(v) for v in row)
            if sheet_terms and not any(term in title_text or term in row_text for term in sheet_terms):
                continue
            if all(norm(needle) in row_text for needle in needles):
                yield title, row

    def feature_statuses(self, feature: str) -> set[str]:
        statuses = set()
        for title, row in self.matching_rows(feature):
            text = norm(title) + " | " + " | ".join(norm(v) for v in row)
            if any(word in text for word in DROP_WORDS):
                statuses.add("drop")
            if any(word in text for word in KEEP_WORDS) or any(word in norm(title) for word in ("final", "selected", "retained", "最终", "入模", "保留")):
                statuses.add("keep")
        return statuses

    def explicit_final_features(self) -> set[str]:
        result = set()
        for title, row in self.rows:
            text = " | ".join(norm(v) for v in row)
            title_final = any(word in norm(title) for word in ("final", "selected", "retained", "最终", "入模", "保留"))
            row_final = any(word in text for word in ("final variable", "final feature", "selected variable", "selected feature", "最终变量", "入模变量"))
            if title_final or row_final:
                result.update(m.group(0).lower() for m in FEATURE_RE.finditer(text))
        return result


def open_view() -> WorkbookView:
    if not OUTPUT.is_file():
        return WorkbookView(f"missing workbook: {OUTPUT}", [])
    try:
        wb = load_workbook(OUTPUT, data_only=False, read_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = [tuple(row) for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)]
            sheets.append((ws.title, rows))
        return WorkbookView(None, sheets)
    except Exception as exc:
        return WorkbookView(f"workbook cannot be opened: {exc}", [])


VIEW = open_view()


def source_truth():
    policy = json.loads((DATA / "screening_policy.json").read_text(encoding="utf-8"))
    raw = pd.read_csv(DATA / "preloan_applications.csv")
    catalog = pd.read_csv(DATA / "feature_catalog.csv")
    iv = pd.read_csv(DATA / "iv_by_org.csv")
    psi = pd.read_csv(DATA / "psi_by_org_transition.csv")
    null = pd.read_csv(DATA / "null_importance.csv")
    valid = raw[raw.target.isin([0, 1])].drop_duplicates("record_id", keep="first").copy()
    valid["month"] = valid.apply_date.astype(str).str.replace("-", "", regex=False).str[:6]
    oos = set(policy["dataset"]["oos_organizations"])
    model = valid[~valid.org_info.isin(oos)].copy()
    month = model.groupby("month").target.agg(total="count", bad="sum").reset_index()
    abnormal = set(month.loc[(month.total < 350) | (month.bad < 20), "month"])
    base = model[~model.month.isin(abnormal)].copy()
    features = catalog.feature.tolist()
    base[features] = base[features].replace([-1, -999, -1111], np.nan)
    return policy, raw, valid, model, base, month, catalog, iv, psi, null


TRUTH = source_truth()


def close(value: float, expected: float, tol: float = 0.0006) -> bool:
    return abs(value - expected) <= tol


def rows_with_number(rows, expected: float, tol: float = 0.0006) -> bool:
    return any(any(close(number, expected, tol) for number in numeric_values(row)) for _, row in rows)


def test_sample_scope_and_exclusions():
    if VIEW.error:
        return
    policy, raw, valid, model, base, month, *_ = TRUTH
    oos_count = int(valid.org_info.isin(policy["dataset"]["oos_organizations"]).sum())
    expected_counts = [
        (len(raw), ("raw", "input", "source records", "applications")),
        (len(valid), ("valid", "dedup", "unique records", "eligible labels", "cleaned rows")),
        (oos_count, ("oos", "out-of-sample", "holdout")),
        (len(base), ("screen", "model", "analysis base", "final sample", "eligible rows")),
    ]
    all_rows = list(VIEW.rows)
    for expected, aliases in expected_counts:
        relevant = [(t, r) for t, r in all_rows if any(a in (norm(t) + " " + " ".join(norm(v) for v in r)) for a in aliases)]
        assert rows_with_number(relevant, expected, tol=0), f"sample flow does not expose the expected {expected} count with a recognizable scope label"

    abnormal_rows = month[(month.total < 350) | (month.bad < 20)]
    for record in abnormal_rows.itertuples(index=False):
        month_id, total, bad = str(record.month), int(record.total), int(record.bad)
        matches = list(VIEW.matching_rows(month_id))
        assert matches, f"abnormal month {month_id} is missing from the workbook"
        reconciled_rows = []
        for title, row in matches:
            numbers = numeric_values(row)
            text = norm(title) + " " + " ".join(norm(v) for v in row)
            if any(close(n, total, 0) for n in numbers) and any(close(n, bad, 0) for n in numbers):
                reconciled_rows.append(text)
        assert reconciled_rows, f"{month_id} sample/bad counts are not reconciled on one reviewable row"
        assert any(any(word in text for word in DROP_WORDS + ("abnormal", "异常")) for text in reconciled_rows), f"{month_id} is not clearly marked as excluded on its quality row"

    expected_pairs = set()
    for org, group in valid.groupby("org_info"):
        for month_id in sorted(group.month.unique()):
            expected_pairs.add((org, month_id))
    represented = set()
    for title, row in all_rows:
        text = " | ".join(norm(v) for v in row)
        for org, month_id in expected_pairs:
            if org.lower() in text and month_id in text:
                represented.add((org, month_id))
    missing_pairs = sorted(expected_pairs - represented)
    assert not missing_pairs, f"organization/month quality coverage is incomplete; missing examples: {missing_pairs[:5]}"


def test_feature_evidence_missingness():
    if VIEW.error:
        return
    *_, base, month, catalog, iv, psi, null = TRUTH
    for feature in ("i_optional_social_score", "i_sparse_device_signal", "i_bank_balance_avg", "i_income_monthly"):
        expected = float(base[feature].isna().mean())
        matches = list(VIEW.matching_rows(feature, sheet_terms=("miss", "缺失")))
        assert matches and rows_with_number(matches, expected), f"missingness evidence for {feature} should include {expected:.4f}"
        structured_values = []
        for title, rows in VIEW.sheets:
            if not rows:
                continue
            headers = [norm(v) for v in rows[0]]
            overall_cols = [i for i, h in enumerate(headers) if ("missing" in h or "缺失" in h) and ("overall" in h or "total" in h or "portfolio" in h or "整体" in h)]
            if not overall_cols:
                continue
            for row in rows[1:]:
                if feature in " | ".join(norm(v) for v in row):
                    for idx in overall_cols:
                        if idx < len(row) and isinstance(row[idx], (int, float)):
                            structured_values.append(float(row[idx]))
        if structured_values:
            assert all(close(value, expected) for value in structured_values), f"conflicting overall missingness values for {feature}: {structured_values}"


def test_feature_evidence_iv_and_psi():
    if VIEW.error:
        return
    iv, psi, null = TRUTH[-3:]
    expected_iv = {"i_low_iv_random": .04, "i_bank_balance_avg": .09, "i_sparse_device_signal": .10, "i_org_weak_signal": .15}
    for feature, expected in expected_iv.items():
        matches = list(VIEW.matching_rows(feature, sheet_terms=("iv", "information value", "信息值")))
        assert matches and rows_with_number(matches, expected), f"overall IV evidence for {feature} should include {expected:.4f}"
    for feature, expected_count in {"i_drifting_utilization": 4, "i_sparse_device_signal": 0}.items():
        matches = list(VIEW.matching_rows(feature, sheet_terms=("psi summary", "psi decision", "psi 处理", "unstable organizations", "不稳定机构")))
        assert matches and rows_with_number(matches, expected_count, tol=0), f"PSI summary for {feature} should show {expected_count} unstable modeling organizations"


def test_feature_evidence_null_importance():
    if VIEW.error:
        return
    for feature, expected in {"i_credit_age_months": 45, "i_noise_hash": 25, "i_sparse_device_signal": 55}.items():
        matches = list(VIEW.matching_rows(feature, sheet_terms=("null", "permut", "noise", "置换", "噪声")))
        assert matches and rows_with_number(matches, expected, tol=0), f"null-importance evidence for {feature} should include gain difference {expected}"


def test_feature_evidence_correlations():
    if VIEW.error:
        return
    base = TRUTH[4]
    pairs = [
        ("i_income_monthly", "i_income_shadow"),
        ("i_debt_ratio", "i_debt_shadow"),
        ("i_loan_to_income", "i_lti_shadow"),
    ]
    for left, right in pairs:
        expected = float(base[[left, right]].corr().abs().iloc[0, 1])
        matches = list(VIEW.matching_rows(left, right, sheet_terms=("corr", "correlation", "相关")))
        assert matches and rows_with_number(matches, expected, tol=.001), f"correlation evidence for {left}/{right} should reconcile to {expected:.4f}"


@pytest.mark.parametrize(
    "gate,expected",
    [
        ("zero_variance", {"i_constant_channel"}),
        ("missingness", {"i_optional_social_score"}),
        ("iv", {"i_bank_balance_avg", "i_low_iv_random", "i_org_weak_signal"}),
        ("psi", {"i_drifting_utilization"}),
        ("null_importance", {"i_credit_age_months", "i_noise_hash"}),
        ("correlation", {"i_debt_shadow", "i_lti_shadow"}),
    ],
)
def test_drop_decisions_by_gate(gate, expected):
    if VIEW.error:
        return
    for feature in expected:
        statuses = VIEW.feature_statuses(feature)
        assert "drop" in statuses, f"{feature} is not explicitly rejected"
        decisive_rows = []
        for title, row in VIEW.matching_rows(feature):
            text = norm(title) + " | " + " | ".join(norm(v) for v in row)
            if any(word in text for word in DROP_WORDS):
                decisive_rows.append(text)
        assert any(any(alias in text for alias in GATE_WORDS[gate]) for text in decisive_rows), f"{feature} does not have a clear {gate} rejection reason"


def test_final_variable_set():
    if VIEW.error:
        return
    expected = {
        "i_income_monthly", "i_income_shadow", "i_debt_ratio", "i_loan_to_income",
        "i_recent_delinquencies", "i_sparse_device_signal", "i_stable_capacity",
    }
    actual = VIEW.explicit_final_features()
    assert actual == expected, f"final modeling variables differ: missing={sorted(expected-actual)}, extra={sorted(actual-expected)}"
