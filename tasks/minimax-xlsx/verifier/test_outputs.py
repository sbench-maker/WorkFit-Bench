from __future__ import annotations

import math
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl
import pytest


OUTPUT = Path("/root/results/trade_settlement_final.xlsx")
INPUT = Path("/root/data/trade_settlement_input.xlsx")
EXPECTED_SHEETS = ["Trades", "Summary", "Fee Schedule", "FX Rates", "Read Me"]
RAW_FIELDS = [
    "trade_id", "trade_date", "desk", "asset_type", "instrument", "currency",
    "side", "status", "quantity", "price", "broker",
]


ALIASES = {
    "trade_id": {"tradeid", "id"},
    "trade_date": {"tradedate", "date"},
    "desk": {"desk", "tradingdesk"},
    "asset_type": {"assettype", "assetclass"},
    "instrument": {"instrument", "security", "ticker"},
    "currency": {"currency", "ccy"},
    "side": {"side", "buysell"},
    "status": {"status", "tradestatus"},
    "quantity": {"quantity", "qty"},
    "price": {"price", "tradeprice"},
    "broker": {"broker", "executingbroker"},
    "gross": {"localgross", "grosslocal", "localnotional", "grossamountlocal"},
    "fee": {"brokerfeeusd", "feeusd", "brokerfee"},
    "cash": {"signedusdcashimpact", "usdcashimpact", "cashimpactusd", "signedcashusd"},
    "active_count": {"activetrades", "activetradecount", "tradecount", "settledtrades"},
}


@dataclass
class Context:
    output_formula: object | None
    input_formula: object | None
    error: str | None


def normalize_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def canonical_header(value) -> str | None:
    normalized = normalize_header(value)
    for field, choices in ALIASES.items():
        if normalized in choices:
            return field
    return None


def normalize_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and math.isfinite(value):
        return round(value, 9)
    return value


def approx(actual, expected, tolerance=0.02) -> bool:
    try:
        return math.isclose(float(actual), float(expected), rel_tol=1e-9, abs_tol=tolerance)
    except (TypeError, ValueError):
        return False


def is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def is_missing_marker(value) -> bool:
    if not isinstance(value, str):
        return False
    normalized = re.sub(r"[^a-z0-9#]", "", value.casefold())
    return (
        "missing" in normalized
        or "pending" in normalized
        or "unavailable" in normalized
        or normalized in {"na", "#na", "#n/a"}
    )


def locate_header(ws, required: set[str], max_rows=15) -> tuple[int, dict[str, int]]:
    for row in range(1, min(ws.max_row, max_rows) + 1):
        columns = {}
        for cell in ws[row]:
            canonical = canonical_header(cell.value)
            if canonical and canonical not in columns:
                columns[canonical] = cell.column
        if required.issubset(columns):
            return row, columns
    raise ValueError(f"{ws.title!r} has no recognizable header row containing {sorted(required)}")


def records_by_id(ws, *, include_derived: bool) -> tuple[dict[str, dict], dict[str, int], int]:
    required = set(RAW_FIELDS)
    if include_derived:
        required.update({"gross", "fee", "cash"})
    header_row, columns = locate_header(ws, required, max_rows=10)
    records = {}
    duplicate_ids = []
    for row in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row, columns["trade_id"]).value
        if value in (None, ""):
            continue
        trade_id = str(value).strip()
        if trade_id in records:
            duplicate_ids.append(trade_id)
            continue
        record = {field: normalize_value(ws.cell(row, columns[field]).value) for field in RAW_FIELDS}
        if include_derived:
            record.update({field: ws.cell(row, columns[field]).value for field in ("gross", "fee", "cash")})
        record["_row"] = row
        records[trade_id] = record
    if duplicate_ids:
        raise ValueError(f"duplicate trade IDs found: {duplicate_ids[:8]}")
    return records, columns, header_row


def sheet_value_matrix(ws, *, max_row=None, max_column=None):
    max_row = ws.max_row if max_row is None else max_row
    max_column = ws.max_column if max_column is None else max_column
    return [
        [normalize_value(ws.cell(row, col).value) for col in range(1, max_column + 1)]
        for row in range(1, max_row + 1)
    ]


@pytest.fixture(scope="session")
def context() -> Context:
    if not OUTPUT.is_file():
        return Context(None, None, f"missing requested workbook: {OUTPUT}")
    try:
        if OUTPUT.stat().st_size < 15_000:
            raise ValueError("output is too small to contain the requested workbook")
        with zipfile.ZipFile(OUTPUT) as archive:
            bad_member = archive.testzip()
            if bad_member:
                raise ValueError(f"corrupt ZIP member: {bad_member}")
        output_formula = openpyxl.load_workbook(OUTPUT, data_only=False, read_only=False)
        input_formula = openpyxl.load_workbook(INPUT, data_only=False, read_only=False)
        return Context(output_formula, input_formula, None)
    except Exception as exc:
        return Context(None, None, f"output workbook is unreadable: {exc}")


def require_context(context: Context):
    if context.error:
        pytest.fail(context.error)
    assert context.output_formula is not None and context.input_formula is not None
    return context.output_formula, context.input_formula


def force_full_recalculation(path: Path) -> None:
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rewritten = path.with_suffix(".force-recalc.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/workbook.xml":
                root = ET.fromstring(payload)
                calc = root.find(f"{{{ns}}}calcPr")
                if calc is None:
                    calc = ET.SubElement(root, f"{{{ns}}}calcPr")
                calc.attrib.update({"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"})
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                root = ET.fromstring(payload)
                changed = False
                for cell in root.iter(f"{{{ns}}}c"):
                    if cell.find(f"{{{ns}}}f") is None:
                        continue
                    cached = cell.find(f"{{{ns}}}v")
                    if cached is not None:
                        cell.remove(cached)
                        changed = True
                if changed:
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, payload)
    os.replace(rewritten, path)


@pytest.fixture(scope="session")
def recalculated(context: Context):
    if context.error:
        return None, context.error
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        return None, "LibreOffice is unavailable despite the environment contract"
    tmp = Path(tempfile.mkdtemp(prefix="settlement-recalc-"))
    source = tmp / "submission.xlsx"
    outdir = tmp / "out"
    profile = tmp / "profile"
    home = tmp / "home"
    outdir.mkdir()
    profile.mkdir()
    home.mkdir()
    shutil.copy2(OUTPUT, source)
    force_full_recalculation(source)
    env = os.environ.copy()
    env["HOME"] = str(home)
    command = [
        office,
        f"-env:UserInstallation=file://{profile}",
        "--headless",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(outdir),
        str(source),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=120, env=env)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return None, f"LibreOffice recalculation failed: {exc}"
    recalculated_path = outdir / "submission.xlsx"
    if result.returncode != 0 or not recalculated_path.is_file():
        detail = (result.stdout + "\n" + result.stderr).strip()[-1200:]
        return None, f"LibreOffice recalculation failed (exit {result.returncode}): {detail}"
    try:
        return openpyxl.load_workbook(recalculated_path, data_only=True, read_only=False), None
    except Exception as exc:
        return None, f"recalculated workbook is unreadable: {exc}"


def source_maps(input_wb):
    source_records, _, _ = records_by_id(input_wb["Trades"], include_derived=False)
    fee_ws = input_wb["Fee Schedule"]
    fee_map = {
        (str(fee_ws.cell(row, 1).value), str(fee_ws.cell(row, 2).value)): float(fee_ws.cell(row, 3).value)
        for row in range(2, fee_ws.max_row + 1)
        if fee_ws.cell(row, 1).value not in (None, "")
    }
    fx_ws = input_wb["FX Rates"]
    fx_map = {}
    for row in range(2, fx_ws.max_row + 1):
        currency = fx_ws.cell(row, 1).value
        if currency in (None, ""):
            continue
        raw_rate = fx_ws.cell(row, 2).value
        fx_map[str(currency)] = None if raw_rate in (None, "") else float(raw_rate)
    return source_records, fee_map, fx_map


def expected_trade_values(record, fee_map, fx_map):
    gross = float(record["quantity"]) * float(record["price"])
    status = str(record["status"]).casefold()
    if status == "cancelled":
        return gross, 0.0, 0.0
    rate = fx_map.get(str(record["currency"]))
    if rate is None:
        return gross, "missing", "missing"
    bps = fee_map[(str(record["broker"]), str(record["asset_type"]))]
    fee = gross * bps / 10000.0 * rate
    gross_usd = gross * rate
    if str(record["side"]).casefold() == "buy":
        cash = -(gross_usd + fee)
    else:
        cash = gross_usd - fee
    return gross, fee, cash


def compare_expected(actual, expected) -> bool:
    if expected == "missing":
        return is_missing_marker(actual)
    return approx(actual, expected)


def summary_rows(ws, header_row: int, columns: dict[str, int]):
    groups = {}
    grand = None
    for row in range(header_row + 1, ws.max_row + 1):
        desk_value = ws.cell(row, columns["desk"]).value
        currency_value = ws.cell(row, columns["currency"]).value
        if desk_value in (None, "") and currency_value in (None, ""):
            continue
        desk_text = str(desk_value or "").strip()
        if normalize_header(desk_text) in {"grandtotal", "overalltotal", "total"}:
            grand = {
                "count": ws.cell(row, columns["active_count"]).value,
                "fee": ws.cell(row, columns["fee"]).value,
                "cash": ws.cell(row, columns["cash"]).value,
                "_row": row,
            }
            continue
        if desk_value in (None, "") or currency_value in (None, ""):
            continue
        key = (str(desk_value).strip(), str(currency_value).strip().upper())
        if key in groups:
            raise ValueError(f"duplicate Summary row for {key}")
        groups[key] = {
            "count": ws.cell(row, columns["active_count"]).value,
            "fee": ws.cell(row, columns["fee"]).value,
            "cash": ws.cell(row, columns["cash"]).value,
            "_row": row,
        }
    return groups, grand


def expected_summary(source_records, fee_map, fx_map):
    grouped = defaultdict(lambda: {"count": 0, "fee": 0.0, "cash": 0.0, "missing": False})
    for record in source_records.values():
        key = (str(record["desk"]), str(record["currency"]).upper())
        gross, fee, cash = expected_trade_values(record, fee_map, fx_map)
        if str(record["status"]).casefold() != "cancelled":
            grouped[key]["count"] += 1
        if fee == "missing" or cash == "missing":
            grouped[key]["missing"] = True
        else:
            grouped[key]["fee"] += float(fee)
            grouped[key]["cash"] += float(cash)
    return dict(grouped)


def test_workbook_integrity_and_source_preservation(context: Context):
    output_wb, input_wb = require_context(context)
    assert set(EXPECTED_SHEETS).issubset(output_wb.sheetnames), (
        f"one or more source sheets are missing: got {output_wb.sheetnames}"
    )
    source_records, _, _ = records_by_id(input_wb["Trades"], include_derived=False)
    output_records, _, _ = records_by_id(output_wb["Trades"], include_derived=True)
    assert len(source_records) == 240, "frozen input should contain 240 trades"
    assert set(output_records) == set(source_records), "trade IDs were added, dropped, or duplicated"
    altered = []
    for trade_id, expected in source_records.items():
        actual = output_records[trade_id]
        differences = [field for field in RAW_FIELDS if actual[field] != expected[field]]
        if differences:
            altered.append((trade_id, differences))
    assert not altered, f"source trade fields were altered for {altered[:8]}"
    for sheet_name in ("Fee Schedule", "FX Rates", "Read Me"):
        source_sheet = input_wb[sheet_name]
        assert sheet_value_matrix(
            output_wb[sheet_name], max_row=source_sheet.max_row, max_column=source_sheet.max_column
        ) == sheet_value_matrix(source_sheet), (
            f"{sheet_name} values changed even though it is a preserved reference sheet"
        )


def test_trade_formula_coverage_and_recalculated_results(context: Context, recalculated):
    output_wb, input_wb = require_context(context)
    recalculated_wb, recalc_error = recalculated
    assert recalc_error is None, recalc_error
    source_records, fee_map, fx_map = source_maps(input_wb)
    formula_records, _, _ = records_by_id(output_wb["Trades"], include_derived=True)
    value_records, _, _ = records_by_id(recalculated_wb["Trades"], include_derived=True)
    assert set(formula_records) == set(source_records), "derived fields do not cover the full trade population"
    missing_formulas = [
        (trade_id, field)
        for trade_id, record in formula_records.items()
        for field in ("gross", "fee", "cash")
        if not is_formula(record[field])
    ]
    assert not missing_formulas, (
        f"calculated cells must remain live formulas; missing formula examples: {missing_formulas[:12]}"
    )
    mismatches = []
    for trade_id, source in source_records.items():
        expected = expected_trade_values(source, fee_map, fx_map)
        actual_record = value_records[trade_id]
        for field, expected_value in zip(("gross", "fee", "cash"), expected):
            if not compare_expected(actual_record[field], expected_value):
                mismatches.append((trade_id, field, actual_record[field], expected_value))
    assert not mismatches, f"recalculated trade results are wrong; examples: {mismatches[:12]}"
    missing_rate_ids = [
        trade_id for trade_id, source in source_records.items()
        if fx_map.get(str(source["currency"])) is None and str(source["status"]).casefold() != "cancelled"
    ]
    assert missing_rate_ids, "fixture must exercise booked trades with a missing FX rate"


def test_summary_coverage_formulas_and_reconciliation(context: Context, recalculated):
    output_wb, input_wb = require_context(context)
    recalculated_wb, recalc_error = recalculated
    assert recalc_error is None, recalc_error
    source_records, fee_map, fx_map = source_maps(input_wb)
    expected = expected_summary(source_records, fee_map, fx_map)

    required = {"desk", "currency", "active_count", "fee", "cash"}
    formula_header_row, formula_columns = locate_header(output_wb["Summary"], required)
    value_header_row, value_columns = locate_header(recalculated_wb["Summary"], required)
    formula_groups, formula_grand = summary_rows(output_wb["Summary"], formula_header_row, formula_columns)
    value_groups, value_grand = summary_rows(recalculated_wb["Summary"], value_header_row, value_columns)

    assert set(formula_groups) == set(expected), (
        f"Summary desk/currency coverage differs from the supplied rows: missing={sorted(set(expected)-set(formula_groups))}, "
        f"extra={sorted(set(formula_groups)-set(expected))}"
    )
    assert formula_grand is not None, "Summary is missing a recognizable grand-total row"
    missing_formulas = []
    for key, record in formula_groups.items():
        for field in ("count", "fee", "cash"):
            if not is_formula(record[field]):
                missing_formulas.append((key, field))
    for field in ("count", "fee", "cash"):
        if not is_formula(formula_grand[field]):
            missing_formulas.append(("Grand Total", field))
    assert not missing_formulas, f"Summary metrics must remain live formulas: {missing_formulas[:12]}"

    assert set(value_groups) == set(expected), "recalculated Summary lost a required desk/currency row"
    mismatches = []
    for key, expected_row in expected.items():
        actual = value_groups[key]
        if int(actual["count"]) != expected_row["count"]:
            mismatches.append((key, "count", actual["count"], expected_row["count"]))
        for field in ("fee", "cash"):
            target = "missing" if expected_row["missing"] else expected_row[field]
            if not compare_expected(actual[field], target):
                mismatches.append((key, field, actual[field], target))
    assert not mismatches, f"Summary values do not reconcile to Trades; examples: {mismatches[:12]}"

    assert value_grand is not None, "recalculated workbook has no grand total"
    assert int(value_grand["count"]) == sum(row["count"] for row in expected.values()), (
        "grand-total trade count does not equal the desk/currency counts"
    )
    assert is_missing_marker(value_grand["fee"]) and is_missing_marker(value_grand["cash"]), (
        "grand total must visibly remain unresolved while booked CZK trades lack an FX rate"
    )
